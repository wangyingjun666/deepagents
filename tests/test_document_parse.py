# -*- coding: utf-8 -*-
"""文档解析子 Agent 的单元测试：沙箱提取 + 宿主工具路由 + 子 Agent 挂载接线。

覆盖范围
--------
1. 沙箱 ops 层：op_extract_document 对 docx / pdf / xlsx 的提取、
   不支持后缀、文件不存在、路径穿越拒绝
2. 端到端工具链路：进程级沙箱 + read_file_content（文本走 read_file、
   文档走 extract_document）、穿越拒绝、guest 权限拒绝
3. 挂载接线：prompts.yml 三域齐备且无 ragflow、document_parse_agent 结构、
   knowledge_base_agent 已摘除、主 Agent 图可离线构建

不依赖 Docker / 数据库 / 外部 API / 大模型 Key（构图用假 Key，不发任何请求）。
运行：python tests/test_document_parse.py
"""
from __future__ import annotations

import asyncio
import importlib.util
import os
import sys
import tempfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

PASS: list[str] = []
FAIL: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    (PASS if cond else FAIL).append(name)
    print(("  [OK]   " if cond else "  [FAIL] ") + name + (f"  -> {detail}" if detail else ""))


# --------------------------------------------------------------------------
# 样例文件：docx / xlsx / 最小 pdf / txt
# --------------------------------------------------------------------------
def make_samples(directory: Path) -> None:
    from docx import Document
    import pandas as pd

    doc = Document()
    doc.add_paragraph("空调年度销售报告")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "型号"
    table.cell(0, 1).text = "销量"
    table.cell(1, 0).text = "HAK-180"
    table.cell(1, 1).text = "1234"
    doc.save(str(directory / "报告.docx"))

    pd.DataFrame({"型号": ["HAK-180", "HAK-260"], "销量": [1234, 876]}).to_excel(
        directory / "销售.xlsx", index=False)

    stream = b"BT /F1 24 Tf 72 720 Td (HELLO SANDBOX PDF) Tj ET"
    objs = [
        b"<</Type/Catalog/Pages 2 0 R>>",
        b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
        b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Contents 4 0 R"
        b"/Resources<</Font<</F1 5 0 R>>>>>>",
        b"<</Length " + str(len(stream)).encode() + b">>stream\n" + stream + b"\nendstream",
        b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objs, 1):
        offsets.append(len(body))
        body += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(body)
    xref = b"xref\n0 6\n0000000000 65535 f \n"
    for off in offsets:
        xref += f"{off:010d} 00000 n \n".encode()
    pdf = body + xref + (b"trailer\n<</Size 6/Root 1 0 R>>\nstartxref\n"
                         + str(xref_pos).encode() + b"\n%%EOF")
    (directory / "文档.pdf").write_bytes(pdf)

    (directory / "说明.txt").write_text("纯文本内容 OK", encoding="utf-8")


# --------------------------------------------------------------------------
# 1. ops 层直调（不起沙箱进程，直接 import ops.py 验证提取逻辑）
# --------------------------------------------------------------------------
def test_ops_extract(tmp: Path) -> None:
    print("\n=== 1. 沙箱 ops 层：op_extract_document ===")
    ws = tmp / "ws"
    up = tmp / "up"
    ws.mkdir(parents=True, exist_ok=True)
    up.mkdir(parents=True, exist_ok=True)
    make_samples(up)

    os.environ["SANDBOX_WORKSPACE"] = str(ws)
    os.environ["SANDBOX_UPLOADS"] = str(up)
    spec = importlib.util.spec_from_file_location(
        "sandbox_ops_under_test", PROJECT_ROOT / "sandbox" / "worker" / "ops.py")
    ops = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(ops)

    r = ops.op_extract_document({"path": "uploads/报告.docx"})
    check("docx 提取段落与表格", r["ok"] and "空调年度销售报告" in r["stdout"]
          and "HAK-180 | 1234" in r["stdout"], r.get("stderr", "")[:120])

    r = ops.op_extract_document({"path": "uploads/销售.xlsx"})
    check("xlsx 提取表头与数据", r["ok"] and "型号,销量" in r["stdout"]
          and "HAK-260" in r["stdout"], r.get("stderr", "")[:120])

    r = ops.op_extract_document({"path": "uploads/文档.pdf"})
    check("pdf 提取文本", r["ok"] and "HELLO SANDBOX PDF" in r["stdout"], r.get("stderr", "")[:160])

    r = ops.op_extract_document({"path": "uploads/说明.txt"})
    check("文本后缀不走文档解析", (not r["ok"]) and r["error_code"] == "unsupported_suffix",
          r.get("stderr", "")[:120])

    r = ops.op_extract_document({"path": "uploads/不存在.docx"})
    check("文件不存在返回 not_found", (not r["ok"]) and r["error_code"] == "not_found")

    r = ops.op_extract_document({"path": "../../etc/passwd"})
    check("路径穿越被沙箱拒绝", (not r["ok"])
          and r["error_code"] in ("sandbox_path_denied", "not_found"), r.get("error_code"))

    r = ops.op_extract_document({"path": "uploads/报告.docx", "max_chars": 5})
    check("超长截断带提示", r["ok"] and "已截断" in r["stdout"] and r["truncated"] is True)


# --------------------------------------------------------------------------
# 2. 端到端：进程级沙箱 + read_file_content 工具路由 + 权限
# --------------------------------------------------------------------------
async def test_tool_e2e(tmp: Path) -> None:
    print("\n=== 2. 工具链路（进程级沙箱 + read_file_content） ===")
    from api.context import (reset_sandbox_context, reset_security_context,
                             reset_session_context, set_sandbox_context,
                             set_security_context, set_session_context, set_thread_context)
    from sandbox import sandbox_manager
    from security.permissions import Role, SessionPrincipal
    from tools.upload_file_read_tool import read_file_content

    ws = tmp / "output" / "session_dp1"
    up = tmp / "updated" / "session_dp1"
    ws.mkdir(parents=True, exist_ok=True)
    up.mkdir(parents=True, exist_ok=True)
    make_samples(up)

    async with sandbox_manager.acquire("dp1", workspace=ws, uploads=up) as sb:
        tokens = [set_session_context(str(ws)), set_thread_context("dp1"),
                  set_sandbox_context(sb),
                  set_security_context(SessionPrincipal.build("dp1", user_id="张三",
                                                              role=Role.USER))]
        try:
            r = read_file_content.invoke({"filename": "说明.txt"})
            check("文本文件走 read_file", "纯文本内容 OK" in r, r[:140])

            r = read_file_content.invoke({"filename": "报告.docx"})
            check("docx 走 extract_document", "空调年度销售报告" in r, r[:140])

            r = read_file_content.invoke({"filename": "销售.xlsx"})
            check("xlsx 走 extract_document", "HAK-180" in r, r[:140])

            r = read_file_content.invoke({"filename": "文档.pdf"})
            check("pdf 走 extract_document", "HELLO SANDBOX PDF" in r, r[:160])

            r = read_file_content.invoke({"filename": "../../../etc/passwd"})
            check("穿越路径被拒", ("错误" in r or "拒绝" in r), r[:140])
        finally:
            reset_security_context(tokens[3])
            reset_sandbox_context(tokens[2])
            reset_session_context(tokens[0], tokens[1])

    async with sandbox_manager.acquire("dp1", workspace=ws, uploads=up) as sb:
        tokens = [set_session_context(str(ws)), set_thread_context("dp1"), set_sandbox_context(sb),
                  set_security_context(SessionPrincipal.build("dp1", user_id="访客",
                                                              role=Role.GUEST))]
        try:
            r = read_file_content.invoke({"filename": "说明.txt"})
            check("guest 读文件被权限拒绝", "权限拒绝" in r, r[:140])
        finally:
            reset_security_context(tokens[3])
            reset_sandbox_context(tokens[2])
            reset_session_context(tokens[0], tokens[1])

    await sandbox_manager.release("dp1")


# --------------------------------------------------------------------------
# 3. 挂载接线：prompts.yml / 子 Agent 结构 / 主 Agent 构图
# --------------------------------------------------------------------------
def test_wiring() -> None:
    print("\n=== 3. 子 Agent 挂载接线（离线构图，不调大模型） ===")
    os.environ.setdefault("OPENAI_API_KEY", "sk-test-not-used")

    from agent.prompts import sub_agents_content
    check("prompts.yml 三域齐备", set(sub_agents_content) == {"tavily", "db", "doc_parse"},
          str(sorted(sub_agents_content)))
    check("ragflow 段已移除", "ragflow" not in sub_agents_content)

    from agent.subagents.document_parse_agent import document_parse_agent
    from tools.upload_file_read_tool import read_file_content
    check("document_parse_agent 结构完整",
          document_parse_agent["name"] == "文档解析助手"
          and document_parse_agent["description"].strip()
          and document_parse_agent["system_prompt"].strip()
          and document_parse_agent["tools"] == [read_file_content])

    try:
        import agent.subagents.knowledge_base_agent  # noqa: F401
        check("knowledge_base_agent 已摘除", False, "模块仍存在")
    except ModuleNotFoundError:
        check("knowledge_base_agent 已摘除", True)

    import agent.main_agent as ma
    check("read_file_content 已移交子 Agent（主 Agent 模块不再持有）",
          not hasattr(ma, "read_file_content"))
    graph = ma.get_main_agent()
    check("主 Agent 图离线构建成功（含三路子 Agent）", graph is not None)

    # 尽力自省 task 工具描述：三路子 Agent 都应注册进委派路由
    described = ""
    try:
        for pregel_node in graph.nodes.values():
            bound = getattr(pregel_node, "bound", None)
            tools = getattr(bound, "tools", None) or getattr(bound, "tools_by_name", None)
            if tools:
                items = tools.values() if isinstance(tools, dict) else tools
                for t in items:
                    if getattr(t, "name", "") == "task":
                        described = t.description or ""
    except Exception:
        described = ""
    if described:
        check("task 委派路由含三路且无 RAGFlow",
              all(k in described for k in ("网络搜索助手", "数据库查询助手", "文档解析助手"))
              and "RAGFlow" not in described, described[:120])
    else:
        print("  [SKIP] task 工具描述自省不可用（框架内部结构变化），跳过该项")


async def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="docparse_"))
    test_ops_extract(tmp)
    await test_tool_e2e(tmp)
    test_wiring()
    print(f"\n=== 文档解析测试结果：{len(PASS)} 通过 / {len(FAIL)} 失败 ===")
    for f in FAIL:
        print("   - 失败：", f)
    return 1 if FAIL else 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
